from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from app.main import create_app

AUTH = {'Authorization': 'Bearer test-access'}


def docx_bytes():
    document = Document()
    document.add_paragraph('请在2026年10月23日前提交项目书。')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = '提交渠道'
    table.cell(0, 1).text = '教学平台'
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_docx_upload_extracts_paragraphs_and_table_without_saving(settings):
    with TestClient(create_app(settings)) as client:
        response = client.post('/api/v1/documents/extract', headers=AUTH,
                               files={'file': ('通知.docx', docx_bytes(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
    assert response.status_code == 200
    assert response.json()['kind'] == 'Word'
    assert response.json()['text'] == '请在2026年10月23日前提交项目书。\n提交渠道 教学平台'
    assert response.json()['characters'] == len(response.json()['text'])


def test_docx_raw_endpoint_supports_app_content_uri_upload(settings):
    with TestClient(create_app(settings)) as client:
        response = client.post('/api/v1/documents/extract-raw?filename=%E9%80%9A%E7%9F%A5.docx', headers={
            **AUTH, 'Content-Type': 'application/octet-stream'
        }, content=docx_bytes())
    assert response.status_code == 200
    assert response.json()['kind'] == 'Word'


def test_document_upload_auth_type_and_size_are_enforced(settings):
    with TestClient(create_app(settings)) as client:
        assert client.post('/api/v1/documents/extract', files={'file': ('通知.docx', docx_bytes())}).status_code == 401
        unsupported = client.post('/api/v1/documents/extract', headers=AUTH, files={'file': ('通知.txt', b'text')})
        assert unsupported.status_code == 415
        too_large = client.post('/api/v1/documents/extract', headers=AUTH, files={'file': ('通知.pdf', b'x' * (12 * 1024 * 1024 + 1))})
        assert too_large.status_code == 413
