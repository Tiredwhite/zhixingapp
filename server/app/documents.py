from io import BytesIO
from pathlib import PurePath
import shutil
import subprocess
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import BadZipFile, ZipFile

from docx import Document
from pypdf import PdfReader

from .errors import ApiError

MAX_DOCUMENT_BYTES = 20 * 1024 * 1024
MAX_EXTRACTED_CHARS = 12000
MAX_PDF_PAGES = 80
MAX_DOCX_ENTRIES = 1000
MAX_DOCX_UNPACKED_BYTES = 64 * 1024 * 1024


def _normalize_text(value: str) -> str:
    lines = [" ".join(line.split()) for line in value.replace("\x00", "").splitlines()]
    return "\n".join(line for line in lines if line)


def _extract_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(data), strict=False)
        if reader.is_encrypted and not reader.decrypt(""):
            raise ApiError(422, "DOCUMENT_ENCRYPTED", "PDF 已加密，暂时无法读取")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ApiError(422, "DOCUMENT_TOO_MANY_PAGES", "PDF 超过 80 页，请拆分后再导入")
        return _normalize_text("\n".join(page.extract_text() or "" for page in reader.pages))
    except ApiError:
        raise
    except Exception as exc:
        raise ApiError(422, "DOCUMENT_INVALID", "PDF 文件损坏或格式不受支持") from exc


def _extract_docx(data: bytes) -> str:
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_DOCX_ENTRIES or sum(entry.file_size for entry in entries) > MAX_DOCX_UNPACKED_BYTES:
                raise ApiError(422, "DOCUMENT_TOO_COMPLEX", "Word 文档结构过大，暂时无法安全读取")
        document = Document(BytesIO(data))
        content = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                content.append("　".join(cell.text for cell in row.cells))
        return _normalize_text("\n".join(content))
    except ApiError:
        raise
    except (BadZipFile, KeyError, ValueError, OSError) as exc:
        raise ApiError(422, "DOCUMENT_INVALID", "Word 文件损坏或格式不受支持") from exc


def _extract_doc(data: bytes) -> str:
    antiword = shutil.which("antiword")
    if not antiword:
        raise ApiError(503, "DOCUMENT_CONVERTER_UNAVAILABLE", "服务器暂未启用旧版 .doc 解析，请先另存为 .docx 再试")
    try:
        with TemporaryDirectory(prefix="zhixing-doc-") as temporary:
            source = Path(temporary) / "notice.doc"
            source.write_bytes(data)
            result = subprocess.run(
                [antiword, "-m", "UTF-8.txt", str(source)],
                stdout=subprocess.PIPE,
                stderr=subprocess.DEVNULL,
                timeout=10,
                check=False,
            )
        if result.returncode != 0:
            raise ApiError(422, "DOCUMENT_INVALID", "Word 文件损坏或格式不受支持")
        return _normalize_text(result.stdout.decode("utf-8", errors="replace"))
    except subprocess.TimeoutExpired as exc:
        raise ApiError(422, "DOCUMENT_TOO_COMPLEX", "旧版 Word 文档处理超时，请另存为 .docx 后再试") from exc
    except ApiError:
        raise
    except OSError as exc:
        raise ApiError(503, "DOCUMENT_CONVERTER_UNAVAILABLE", "服务器暂时无法读取旧版 .doc 文档") from exc
    except Exception as exc:
        raise ApiError(422, "DOCUMENT_INVALID", "Word 文件损坏或格式不受支持") from exc


def extract_document(filename: str | None, data: bytes) -> tuple[str, str]:
    extension = PurePath(filename or "").suffix.lower()
    if len(data) > MAX_DOCUMENT_BYTES:
        raise ApiError(413, "DOCUMENT_TOO_LARGE", "文件不能超过 20 MB")
    if not data:
        raise ApiError(400, "DOCUMENT_EMPTY", "所选文件为空")
    if extension == ".pdf":
        text = _extract_pdf(data)
        kind = "PDF"
    elif extension == ".docx":
        text = _extract_docx(data)
        kind = "Word"
    elif extension == ".doc":
        text = _extract_doc(data)
        kind = "Word"
    else:
        raise ApiError(415, "DOCUMENT_UNSUPPORTED", "目前支持 PDF、.docx 和 .doc Word 文档")
    if not text:
        raise ApiError(422, "DOCUMENT_NO_TEXT", "未提取到可选择的文字，扫描版 PDF 暂不支持 OCR")
    if len(text) > MAX_EXTRACTED_CHARS:
        raise ApiError(413, "DOCUMENT_TEXT_TOO_LONG", "提取文字超过 12000 字，请先拆分或删去无关页面")
    return text, kind
